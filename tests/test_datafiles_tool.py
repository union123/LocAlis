"""
Тесты инструмента создания файлов данных (datafiles).

Раньше этого файла не было вообще — write_excel падал в проде с
AttributeError: 'str' object has no attribute 'items' (задача 6d82e8b850c1,
17.08), и это не было закрыто ни одним тестом. Закрепляем:
  - обычный dict sheets работает;
  - модель передаёт sheets как JSON-строку (и даже дважды закодированную) —
    разворачиваем молча, без падения и без потери данных в "Данные"-заглушку;
  - защита от записи в файл с чужим расширением (типичная порча исходных
    данных моделью, см. комментарий в datafiles_tool.py про задачу 34ce9951);
  - защита от перезаписи без overwrite=true;
  - базовые happy-path для остальных write_* действий.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tools.datafiles.datafiles_tool import DataFilesTool


@pytest.fixture()
def tool() -> DataFilesTool:
    return DataFilesTool(config={"allowed_roots": []})


# ---- write_excel ------------------------------------------------------


def test_write_excel_accepts_dict_sheets(tool: DataFilesTool, tmp_path: Path):
    out = tmp_path / "report.xlsx"
    res = tool.execute("write_excel", {"output_path": str(out),
                                        "sheets": {"S1": [{"a": 1, "b": 2}]}})
    assert res["ok"], res.get("error")
    wb = load_workbook(out)
    assert wb["S1"]["A1"].value == "a"
    assert wb["S1"]["A2"].value == 1


def test_write_excel_accepts_json_encoded_string_sheets(tool: DataFilesTool, tmp_path: Path):
    """Модели часто передают sheets как JSON-строку, а не словарь."""
    out = tmp_path / "report.xlsx"
    sheets_str = json.dumps({"S1": [{"a": 1}]})
    res = tool.execute("write_excel", {"output_path": str(out), "sheets": sheets_str})
    assert res["ok"], res.get("error")
    wb = load_workbook(out)
    assert "S1" in wb.sheetnames


def test_write_excel_accepts_double_json_encoded_sheets(tool: DataFilesTool, tmp_path: Path):
    """Защита от AttributeError: 'str' object has no attribute 'items'
    (наблюдалось в проде, задача 6d82e8b850c1) при двойном JSON-кодировании."""
    out = tmp_path / "report.xlsx"
    sheets_str = json.dumps(json.dumps({"S1": [{"a": 1}]}))
    res = tool.execute("write_excel", {"output_path": str(out), "sheets": sheets_str})
    assert res["ok"], res.get("error")
    wb = load_workbook(out)
    # Реальные данные должны попасть в лист "S1", а не потеряться в
    # текстовой заглушке "Данные" (это была бы тихая порча данных).
    assert "S1" in wb.sheetnames
    assert "Данные" not in wb.sheetnames


def test_write_excel_wrong_extension_is_rejected(tool: DataFilesTool, tmp_path: Path):
    """Модели путают output_path и пишут excel-данные в .geojson — блокируем."""
    out = tmp_path / "report.geojson"
    res = tool.execute("write_excel", {"output_path": str(out),
                                        "sheets": {"S1": [{"a": 1}]}})
    assert not res["ok"]
    assert not out.exists()


def test_write_excel_existing_corrupt_file_still_reported_as_error(tool: DataFilesTool, tmp_path: Path):
    """Не-xlsx мусор по этому пути нельзя бесшумно дописать как книгу."""
    out = tmp_path / "report.xlsx"
    out.write_text("уже существует, но это не xlsx")
    res = tool.execute("write_excel", {"output_path": str(out),
                                        "sheets": {"S1": [{"a": 1}]}})
    assert not res["ok"]
    res_ok = tool.execute("write_excel", {"output_path": str(out),
                                           "sheets": {"S1": [{"a": 1}]},
                                           "overwrite": True})
    assert res_ok["ok"], res_ok.get("error")


def test_write_excel_multiple_calls_append_sheets(tool: DataFilesTool, tmp_path: Path):
    """Сборка большой книги по частям: повторный вызов с overwrite=false
    добавляет лист, а не требует overwrite=true и не теряет уже записанные
    листы — это и даёт маленьким моделям возможность собирать отчёт по частям."""
    out = tmp_path / "report.xlsx"
    res1 = tool.execute("write_excel", {"output_path": str(out),
                                         "sheets": {"Сводка": [{"a": 1}]}})
    assert res1["ok"], res1.get("error")
    res2 = tool.execute("write_excel", {"output_path": str(out),
                                         "sheets": {"Данные": [{"b": 2}]},
                                         "overwrite": False})
    assert res2["ok"], res2.get("error")
    wb = load_workbook(out)
    assert set(wb.sheetnames) == {"Сводка", "Данные"}


def test_write_excel_overwrite_true_discards_previous_sheets(tool: DataFilesTool, tmp_path: Path):
    out = tmp_path / "report.xlsx"
    tool.execute("write_excel", {"output_path": str(out), "sheets": {"Старый": [{"a": 1}]}})
    res = tool.execute("write_excel", {"output_path": str(out),
                                        "sheets": {"Новый": [{"b": 2}]},
                                        "overwrite": True})
    assert res["ok"], res.get("error")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Новый"]


def test_write_excel_malformed_json_gives_actionable_error(tool: DataFilesTool, tmp_path: Path):
    """Синтаксически битый JSON (не дважды закодированный валидный,
    а просто сломанный) должен давать ошибку с позицией и советом, а не
    глухой AttributeError/трассировкой openpyxl."""
    out = tmp_path / "report.xlsx"
    broken = '{"S1": [{"a" 1}]}'  # нет ":" между "a" и 1
    res = tool.execute("write_excel", {"output_path": str(out), "sheets": broken})
    assert not res["ok"]
    assert "позиции" in res["error"]
    assert not out.exists()


# ---- write_geojson / write_csv / write_json / write_text ---------------


def test_write_geojson_wraps_bare_properties(tool: DataFilesTool, tmp_path: Path):
    out = tmp_path / "layer.geojson"
    res = tool.execute("write_geojson", {"output_path": str(out),
                                          "features": [{"name": "x"}]})
    assert res["ok"], res.get("error")
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["type"] == "FeatureCollection"
    assert data["features"][0]["type"] == "Feature"
    assert data["features"][0]["properties"] == {"name": "x"}


def test_write_csv_basic(tool: DataFilesTool, tmp_path: Path):
    out = tmp_path / "table.csv"
    res = tool.execute("write_csv", {"output_path": str(out),
                                      "rows": [{"a": 1, "b": 2}]})
    assert res["ok"], res.get("error")
    text = out.read_text(encoding="utf-8-sig")
    assert "a" in text and "1" in text


def test_write_json_basic(tool: DataFilesTool, tmp_path: Path):
    out = tmp_path / "data.json"
    res = tool.execute("write_json", {"output_path": str(out), "data": {"k": "v"}})
    assert res["ok"], res.get("error")
    assert json.loads(out.read_text(encoding="utf-8")) == {"k": "v"}


def test_write_text_basic(tool: DataFilesTool, tmp_path: Path):
    out = tmp_path / "note.txt"
    res = tool.execute("write_text", {"output_path": str(out), "content": "привет"})
    assert res["ok"], res.get("error")
    assert out.read_text(encoding="utf-8") == "привет"


# ---- относительные пути (задача 4ce3abd8, 21.08) ------------------------


def test_write_text_bare_relative_filename_is_rejected(tool: DataFilesTool):
    """Задача 4ce3abd8: модель передала голые имена файлов ('init-db.ts',
    'server.ts') без пути к папке задачи. normalize() их не трогает (нет ни
    известной папки, ни буквы диска), поэтому раньше Path(...) оставался
    относительным и запись молча уходила в рабочую папку процесса вместо
    нужной папки задачи — без единой ошибки. Теперь это должно быть явной
    ошибкой, а не тихой записью не туда."""
    stray = Path.cwd() / "server.ts"
    assert not stray.exists(), "файл не должен существовать до вызова инструмента"
    res = tool.execute("write_text", {"output_path": "server.ts", "content": "x"})
    assert not res["ok"]
    assert "относительн" in res["error"].lower()
    assert not stray.exists(), "инструмент не должен молча писать в рабочую папку процесса"


def test_write_text_absolute_path_still_works(tool: DataFilesTool, tmp_path: Path):
    """Абсолютный путь (в т.ч. к ещё не созданной вложенной папке) должен
    продолжать работать как раньше — родители создаются автоматически."""
    out = tmp_path / "site_calisthenics" / "backend" / "server.ts"
    res = tool.execute("write_text", {"output_path": str(out), "content": "x"})
    assert res["ok"], res.get("error")
    assert out.read_text(encoding="utf-8") == "x"


# ---- allowed_roots на уровне задачи (21.08, корень проблемы edc0c5668a84) ---
#
# Промпт-уровневый запрет писать в конкретную папку модели дважды игнорировали.
# Теперь это должна быть проверка на уровне инструмента через GET_ALLOWED_ROOTS
# в контексте, независимо от текста в промпте.


def test_task_allowed_roots_blocks_write_outside_project_folder(tmp_path: Path):
    """Главная регрессия этого исправления: даже без статического allowed_roots в
    manifest.yaml, задача с get_allowed_roots() в контексте должна быть жёстко
    ограничена своей папкой проекта."""
    project_dir = tmp_path / "site_calisthenics"
    project_dir.mkdir()
    other_dir = tmp_path / "other_project"
    other_dir.mkdir()
    tool = DataFilesTool(
        config={"allowed_roots": []},
        context={"get_allowed_roots": lambda: [str(project_dir)]},
    )
    bad = other_dir / "server.ts"
    res = tool.execute("write_text", {"output_path": str(bad), "content": "x"})
    assert not res["ok"]
    assert not bad.exists()


def test_task_allowed_roots_permits_write_inside_project_folder(tmp_path: Path):
    project_dir = tmp_path / "site_calisthenics"
    project_dir.mkdir()
    tool = DataFilesTool(
        config={"allowed_roots": []},
        context={"get_allowed_roots": lambda: [str(project_dir)]},
    )
    good = project_dir / "backend" / "server.ts"
    res = tool.execute("write_text", {"output_path": str(good), "content": "x"})
    assert res["ok"], res.get("error")
    assert good.read_text(encoding="utf-8") == "x"


def test_task_allowed_roots_take_priority_over_static_config(tmp_path: Path):
    """Ограничение задачи перекрывает статический allowed_roots из manifest.yaml,
    а не объединяется с ним и не игнорируется."""
    static_root = tmp_path / "static_root"
    static_root.mkdir()
    task_root = tmp_path / "task_root"
    task_root.mkdir()
    tool = DataFilesTool(
        config={"allowed_roots": [str(static_root)]},
        context={"get_allowed_roots": lambda: [str(task_root)]},
    )
    blocked = static_root / "file.txt"
    res_blocked = tool.execute("write_text", {"output_path": str(blocked), "content": "x"})
    assert not res_blocked["ok"]
    assert not blocked.exists()
    allowed = task_root / "file.txt"
    res_allowed = tool.execute("write_text", {"output_path": str(allowed), "content": "x"})
    assert res_allowed["ok"], res_allowed.get("error")


def test_static_allowed_roots_still_enforced_without_task_context(tmp_path: Path):
    """Если задача своё ограничение не задаёт, статический allowed_roots из
    manifest.yaml должен продолжать работать как раньше."""
    static_root = tmp_path / "static_root"
    static_root.mkdir()
    other_dir = tmp_path / "outside"
    other_dir.mkdir()
    tool = DataFilesTool(
        config={"allowed_roots": [str(static_root)]},
        context={"get_allowed_roots": lambda: []},
    )
    blocked = other_dir / "file.txt"
    res = tool.execute("write_text", {"output_path": str(blocked), "content": "x"})
    assert not res["ok"]
    assert not blocked.exists()
    allowed = static_root / "file.txt"
    res2 = tool.execute("write_text", {"output_path": str(allowed), "content": "x"})
    assert res2["ok"], res2.get("error")


def test_multiple_task_allowed_roots_all_checked(tmp_path: Path):
    """Старый баг: _path() проверял только roots[0], игнорируя остальные корни
    в списке. Задача с двумя разрешёнными корнями должна принимать запись в оба."""
    root1 = tmp_path / "root1"
    root1.mkdir()
    root2 = tmp_path / "root2"
    root2.mkdir()
    tool = DataFilesTool(
        config={"allowed_roots": []},
        context={"get_allowed_roots": lambda: [str(root1), str(root2)]},
    )
    for root in (root1, root2):
        target = root / "file.txt"
        res = tool.execute("write_text", {"output_path": str(target), "content": "x"})
        assert res["ok"], res.get("error")
        assert target.read_text(encoding="utf-8") == "x"
